"""Local document utilities. No model calls; documents are always data."""
import csv
import io
import json
import re
import sys
import zipfile
from pathlib import Path


def text(v):
    if v is None:
        return ''
    return str(v)


def template(path):
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in book:
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 20), values_only=True):
            headers = [text(v).strip() for v in row]
            while headers and not headers[-1]:
                headers.pop()
            if any(headers):
                if len(headers) > 100:
                    raise ValueError('表头超过 100 列，请拆分提取任务')
                headers = [h or f'未命名列{i+1}' for i, h in enumerate(headers)]
                if len(set(headers)) != len(headers):
                    raise ValueError('表头存在重复列名，请修改后上传')
                return {'sheet': sheet.title, 'headers': headers}
    raise ValueError('工作簿没有可用表头')


def example_rows(path):
    """Import completed rows as teaching examples, never as extracted evidence."""
    p = Path(path)
    if p.suffix.lower() == '.csv':
        raw = p.read_bytes()
        try:
            data = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            data = raw.decode('gb18030')
        sheets = [('CSV', list(csv.reader(io.StringIO(data))))]
    else:
        import openpyxl
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheets = []
        for s in book:
            if s.max_row > 2001:
                raise ValueError('示范结果单工作表最多2000条数据，请拆分或精简示范案例')
            sheets.append((s.title, list(s.iter_rows(values_only=True))))
    headers = None
    results, warnings = [], []
    for name, rows in sheets:
        head_idx = next((i for i, row in enumerate(rows[:20]) if any(v is not None and str(v).strip() for v in row)), None)
        if head_idx is None:
            continue
        head = [text(v).strip() for v in rows[head_idx]]
        while head and not head[-1]:
            head.pop()
        if not head or len(head) > 100 or len(set(head)) != len(head) or any(not x for x in head):
            if headers:
                warnings.append(f'{name} 表头不同，未作为示范结果导入')
                continue
            raise ValueError('示范表头须唯一且非空，最多100列；请使用第一行作为表头')
        if headers is None:
            headers = head
        if head != headers:
            warnings.append(f'{name} 与首张结果表列名不同，未导入（例如证据、历史工作表）')
            continue
        for i, row in enumerate(rows[head_idx+1:], head_idx+2):
            if not any(v is not None and str(v).strip() for v in row):
                continue
            if any(v is not None and str(v).strip() for v in row[len(headers):]):
                raise ValueError(f'{name} 第{i}行存在无表头的数据')
            values = {}
            for j, h in enumerate(headers):
                value = row[j] if j < len(row) else None
                if value == '': value = None
                if value is not None and not isinstance(value, (str, int, float, bool)):
                    value = str(value)
                values[h] = value
            results.append({'sheet': name, 'row': i, 'values': values})
    if not results:
        raise ValueError('结果表没有已提取的数据行；只有表头请使用试提取流程')
    if len(results) > 2000:
        raise ValueError('示范案例最多2000行，请提供少量有代表性的已提取结果')
    return {'headers': headers, 'rows': results, 'warnings': warnings}


def unpack(path, output):
    dest = Path(output).resolve()
    dest.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path) as z:
        files = [x for x in z.infolist() if not x.is_dir()]
        if len(files) > 500 or sum(x.file_size for x in files) > 500 * 1024 * 1024:
            raise ValueError('压缩包超出限制：最多 500 个文件、解压后 500 MB')
        for item in files:
            target = (dest / item.filename).resolve()
            if not target.is_relative_to(dest) or (item.external_attr >> 16) & 0o170000 == 0o120000:
                raise ValueError('压缩包包含不安全路径或符号链接')
            if item.file_size > 100 * 1024 * 1024:
                raise ValueError('压缩包中单文件超过 100 MB')
        out = []
        for item in files:
            if '__MACOSX' in item.filename or Path(item.filename).name.startswith('.'):
                continue
            target = dest / item.filename
            target.parent.mkdir(parents=True, exist_ok=True)
            with z.open(item) as src, target.open('wb') as dst:
                import shutil
                shutil.copyfileobj(src, dst)
            out.append({'path': str(target), 'name': item.filename, 'size': item.file_size})
        return out


def table_blocks(path):
    p = Path(path)
    blocks = []
    warnings = []
    if p.suffix.lower() == '.csv':
        raw = p.read_bytes()
        try:
            data = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            data = raw.decode('gb18030')
        sheets = [('CSV', list(csv.reader(io.StringIO(data))))]
    else:
        import openpyxl
        w = openpyxl.load_workbook(p, read_only=True, data_only=True)
        sheets = []
        for s in w:
            if s.max_row > 5000:
                warnings.append(f'{s.title} 超过 5000 行，仅解析前 5000 行；请拆分文件后重新提取')
            sheets.append((s.title, list(s.iter_rows(max_row=min(s.max_row, 5000), values_only=True))))
    for sheet, rows in sheets:
        if len(rows) > 5000:
            warnings.append(f'{sheet} 仅解析前 5000 行')
            rows = rows[:5000]
        nonempty = [(i+1, row) for i, row in enumerate(rows) if any(v is not None and str(v) != '' for v in row)]
        if not nonempty:
            continue
        header = ' | '.join(text(v) for v in nonempty[0][1])
        for start in range(0, len(nonempty), 12):
            batch = nonempty[start:start+12]
            body = '\n'.join(f'行 {i}: ' + ' | '.join(text(v) for v in row) for i, row in batch)
            blocks.append({'type': 'table', 'text': f'工作表 {sheet}\n表头: {header}\n{body}', 'sheet': sheet, 'cell_range': f'A{batch[0][0]}:{column(max(len(row) for _, row in batch))}{batch[-1][0]}', 'page_idx': None, 'bbox': None})
    return {'blocks': blocks, 'warnings': warnings, 'page_count': len(sheets)}


def column(n):
    from openpyxl.utils import get_column_letter
    return get_column_letter(max(n, 1))


def plain_blocks(path):
    p = Path(path)
    if p.suffix.lower() == '.docx':
        from docx import Document
        d = Document(p)
        from docx.oxml.ns import qn
        parts = []
        for el in d.element.body:
            if el.tag == qn('w:p'):
                value = ''.join(el.itertext())
                # python-docx exposes text through paragraph wrapper without duplicate XML text.
                from docx.text.paragraph import Paragraph
                value = Paragraph(el, d).text
                if value.strip():
                    parts.append(('text', value))
            elif el.tag == qn('w:tbl'):
                from docx.table import Table
                t = Table(el, d)
                parts.append(('table', '\n'.join(' | '.join(c.text for c in r.cells) for r in t.rows)))
    else:
        content = p.read_text(encoding='utf-8-sig', errors='replace')
        parts = [('text', x) for x in re.split(r'\n\s*\n', content) if x.strip()]
    blocks = []
    for i, (kind, value) in enumerate(parts):
        for offset in range(0, len(value), 12000):
            blocks.append({'type': kind, 'text': value[offset:offset+12000], 'anchor': f'段落 {i+1} · 字符 {offset+1}–{min(offset+12000, len(value))}', 'page_idx': None, 'bbox': None})
    return {'blocks': blocks, 'warnings': [], 'page_count': 0}


def render(path, page, output):
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(path)
    if page < 0 or page >= len(doc):
        raise ValueError('页码超出范围')
    obj = doc[page]
    obj.render(scale=1.7).to_pil().save(output, format='PNG')
    return {'pages': len(doc), 'width': obj.get_width(), 'height': obj.get_height()}


def visual(path, page, bbox, output):
    from PIL import Image
    if Path(path).suffix.lower() == '.pdf':
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(path)
        if page < 0 or page >= len(doc):
            raise ValueError('图像页码超出范围')
        im = doc[page].render(scale=2.3).to_pil().convert('RGB')
    else:
        im = Image.open(path).convert('RGB')
    width, height = im.size
    im = im.crop(tuple(round(v / 1000 * (width if i % 2 == 0 else height)) for i, v in enumerate(bbox)))
    im.thumbnail((1600, 1800))
    im.save(output, format='PNG')
    return {'width': im.width, 'height': im.height}


def grid_lines(path):
    from PIL import Image
    import numpy as np
    im = np.asarray(Image.open(path).convert('RGB')).astype(float)
    gray = (im.max(2) - im.min(2) < 22) & (im.mean(2) < 170)
    def lines(axis):
        score = gray.mean(axis=axis)
        groups = []
        for x in np.where(score > .47)[0]:
            if not groups or x > groups[-1][-1] + 2:
                groups.append([int(x)])
            else:
                groups[-1].append(int(x))
        return [round(float(np.mean(g)) / len(score) * 1000, 1) for g in groups]
    return {'width': int(im.shape[1]), 'height': int(im.shape[0]), 'vertical_lines': lines(0), 'horizontal_lines': lines(1), 'note': '原图像素检测的长直线位置（图片0–1000坐标），可能包含外框，须与可见经纬度标签对应；不是自动识别出的坐标值。'}


def export_book(project, output, base_url):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.comments import Comment
    w = openpyxl.Workbook()
    ws = w.active
    ws.title = '提取结果'
    fields = project['fields']
    paper_column = project.get('experience') == 'extract'
    ws.append((['Paper'] if paper_column else []) + [f['name'] for f in fields])
    evidence = w.create_sheet('证据与复核')
    evidence.append(['记录', '字段', '提取值', '状态', '说明', '文件', 'PDF页序 / 位置', '证据ID', '原文摘录', '复核链接', '人工复核'])
    missing = w.create_sheet('待复核与缺失')
    missing.append(['记录', '字段', '状态', '说明'])
    estimates = w.create_sheet('图片估算依据')
    estimates.append(['记录', '字段', '估算值', '位置含义', '定位说明', '方法', '参照依据', '纬度不确定范围', '经度不确定范围', '不确定性说明（非置信区间）', '图像证据ID', '原页定位（0–1000）', '图像哈希', '运行ID', '允许估算的用户要求', '地名配准信息（GeoNames CC BY 4.0）'])
    blocks = {b['id']: b for b in project.get('blocks', [])}
    files = {f['id']: f for f in project.get('files', [])}
    def safe(v):
        if isinstance(v, str) and v.startswith(('=', '+', '-', '@')):
            return "'" + v
        return v
    for i, record in enumerate(project.get('records', []), 2):
        ws.append(([safe(record['label'])] if paper_column else []) + [safe(record['cells'].get(f['name'], {}).get('value')) for f in fields])
        for j, field in enumerate(fields, 2 if paper_column else 1):
            c = record['cells'].get(field['name'], {})
            note = c.get('note', '')
            refs = c.get('evidence_ids', [])
            ws.cell(i, j).comment = Comment(f"记录: {record['label']}\n状态: {c.get('status', 'not_found')}\n{note}\n证据: {', '.join(refs)}", 'SIE-Agent')
            if c.get('status') == 'estimated':
                e = c['estimate']
                ws.cell(i, j).fill = PatternFill('solid', fgColor='FFF0CD')
                ws.cell(i, j).number_format = '"≈ "0.##'
                if e.get('kind') == 'visual_readout':
                    if '图像数值估读' not in w.sheetnames:
                        readouts = w.create_sheet('图像数值估读')
                        readouts.append(['记录', '字段', '对象 / 分组', '估读值', '单位', '读图范围下限', '读图范围上限', '读图依据', '误差说明（非统计区间）', '图像证据ID', '图像哈希', '运行ID'])
                    readouts = w['图像数值估读']
                    for item in e['items']:
                        readouts.append([safe(record['label']), safe(field['name']), safe(item['label']), item['value'], safe(item['unit']), *item['range'], safe(item['basis']), safe(item['uncertainty_note']), item['evidence_id'], blocks.get(item['evidence_id'], {}).get('image_hash', ''), e['run_id']])
                    ranges = '\n'.join(f"{item['label']}: {item['range']} {item['unit']}" for item in e['items'])
                    ws.cell(i, j).comment = Comment(f"图片估算，非原文报告。\n{note}\n读图范围（非统计区间）:\n{ranges}\n图像证据: {', '.join(refs)}", 'SIE-Agent')
                else:
                    estimates.append([safe(record['label']),safe(field['name']),safe(c['value']),e['scope'],safe(e['scope_note']),e['method'],safe(e['basis']),json.dumps(e['latitude_range']),json.dumps(e['longitude_range']),safe(e['uncertainty_note']),e['evidence_id'],json.dumps(e['image_point']),e.get('source_image_hash',''),e['run_id'],safe(e.get('authorization','')),json.dumps(e.get('calibration'),ensure_ascii=False)])
                    ws.cell(i, j).comment = Comment(f"图片估算，非原文报告。\n{note}\n纬度范围: {e['latitude_range']}\n经度范围: {e['longitude_range']}\n{e['uncertainty_note']}\n图像证据: {e['evidence_id']}", 'SIE-Agent')
            if c.get('review') != 'accepted' and c.get('status') in ['conflict', 'needs_review', 'not_found', 'estimated']:
                missing.append([safe(record['label']), safe(field['name']), c.get('status'), safe(note)])
            for ref in refs or ['']:
                b = blocks.get(ref, {})
                f = files.get(b.get('file_id'), {})
                location = (b['page_idx'] + 1) if b.get('page_idx') is not None else b.get('cell_range') or b.get('anchor', '')
                url = f"{base_url}/?project={project['id']}&record={record['id']}&field={field['id']}"
                evidence.append([safe(record['label']),safe(field['name']),safe(c.get('value')),c.get('status'),safe(note),safe(f.get('name','')),location,ref,safe(b.get('text','')[:32000]),url,c.get('review','pending')])
                evidence.cell(evidence.max_row, 10).hyperlink = url
    history = w.create_sheet('修订历史')
    history.append(['时间', '记录', '字段', '修改前', '修改后', '原因', '运行ID'])
    for h in project.get('history', []):
        history.append([h.get('at'),safe(h.get('record_label','')),safe(h.get('field','')),safe(text((h.get('before') or {}).get('value'))),safe(text((h.get('after') or {}).get('value'))),safe(h.get('reason','')),safe(h.get('run_id',''))])
    meta = w.create_sheet('提取说明')
    meta.append(['项目',safe(project['name'])])
    meta.append(['行粒度',safe(project.get('granularity'))])
    meta.append(['提取方式','按列名和可选列说明自主识别正文、表格及图片证据'])
    models = list(dict.fromkeys(r.get('model') for r in project.get('runs', []) if r.get('model')))
    meta.append(['模型',safe('、'.join(models) if models else '未记录')])
    meta.append(['导出说明','空白代表未获得可提交的值；请同时查看证据与复核、待复核与缺失工作表。生成ID不代表论文原始编号。'])
    meta.append(['经验版本',json.dumps(project.get('applied_memories',[]),ensure_ascii=False)])
    meta.append(['当前坐标估算约定',json.dumps(project.get('coordinate_policy',{}),ensure_ascii=False)])
    workflow = project.get('workflow')
    if workflow:
        meta.append(['提取流程', '案例学习' if workflow['mode'] == 'examples' else '试提取校准'])
        meta.append(['当前已确认规范', workflow.get('active_version')])
        protocols = w.create_sheet('项目提取规范')
        protocols.append(['版本', '状态', '字段', '规则', '依据类型', '案例行ID', '原文证据ID', '保存时间'])
        for version in workflow.get('protocols', []):
            for rule in version['rules']:
                protocols.append([version['version'], safe(version['status']), safe(rule['field']), safe(rule['instruction']), safe(rule['basis']), safe(', '.join(rule.get('example_ids', []))), safe(', '.join(rule.get('evidence_ids', []))), version['created_at']])
        runs = w.create_sheet('提取运行与规范版本')
        runs.append(['运行ID', '操作', '状态', '所用规范版本', '文件ID', '开始时间'])
        for run in project.get('runs', []):
            runs.append([run['id'], safe(run.get('workflow_action') or run['mode']), run['status'], run.get('protocol_version'), safe(', '.join(run.get('input_file_ids', []))), run['started_at']])
    for sheet in w:
        sheet.freeze_panes = 'A2'
        if sheet.max_row > 1:
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(color='FFFFFF',bold=True)
            cell.fill = PatternFill('solid',fgColor='173D41')
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top',wrap_text=True)
        for col in sheet.columns:
            label = col[0].column_letter
            sheet.column_dimensions[label].width = min(48,max(15,len(str(col[0].value or ''))*1.4))
    w.save(output)
    return {'path': output, 'records': len(project.get('records', []))}


if __name__ == '__main__':
    try:
        a = json.load(sys.stdin)
        op = a['op']
        if op == 'template': result = template(a['path'])
        elif op == 'examples': result = example_rows(a['path'])
        elif op == 'unpack': result = unpack(a['path'], a['output'])
        elif op == 'table': result = table_blocks(a['path'])
        elif op == 'plain': result = plain_blocks(a['path'])
        elif op == 'render': result = render(a['path'], int(a['page']), a['output'])
        elif op == 'visual': result = visual(a['path'], int(a['page']), a['bbox'], a['output'])
        elif op == 'grid_lines': result = grid_lines(a['path'])
        elif op == 'export': result = export_book(a['project'], a['output'], a['base_url'])
        else: raise ValueError('Unknown operation')
        print(json.dumps(result,ensure_ascii=False,default=str))
    except Exception as e:
        print(json.dumps({'error':str(e)},ensure_ascii=False))
        sys.exit(1)
